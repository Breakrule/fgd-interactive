import streamlit as st
import pandas as pd
import json
from datetime import datetime
from fpdf import FPDF
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

# Configure page settings
st.set_page_config(
    page_title="FGD I - Geopark Kotabaru Interactive Instrument v3",
    page_icon="🌍",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Initialize Session State for interactive data collection
if "matrix_data" not in st.session_state:
    st.session_state.matrix_data = [
        {
            "No": 1,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Kelembagaan",
            "Nama Program / Kegiatan": "Pembentukan Tim Agile",
            "Rencana Aksi / Deskripsi": "Pembentukan tim koordinasi internal & agile team untuk menyusun pondasi awal rencana aksi.",
            "OPD Lead": "Sekretariat Daerah (Asisten III)",
            "RACI Support": "Bapperida (R), Bagian Organisasi (R), Bagian Hukum (R), BKPSDMD (R)",
            "KPI": "Tersusunnya SK Tim Kerja & pembagian tugas.",
            "Timeline": "Agustus 2026 (Mg 1-2)"
        },
        {
            "No": 2,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Regulasi",
            "Nama Program / Kegiatan": "Penyusunan Draft Keputusan & Perbup GCG",
            "Rencana Aksi / Deskripsi": "Menyusun draf Peraturan Bupati tentang Tata Kelola Kolaboratif Lintas Perangkat Daerah dalam mendukung Geopark Kotabaru.",
            "OPD Lead": "Sekretariat Daerah (Asisten III)",
            "RACI Support": "Bagian Hukum (R), Bapperida (C), Disparpora (C)",
            "KPI": "Tersedianya Draft Perbup GCG & Keputusan Bupati.",
            "Timeline": "Agustus 2026 (Mg 3-4)"
        },
        {
            "No": 3,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "SOP",
            "Nama Program / Kegiatan": "Penyusunan SOP Tata Kelola Kolaboratif",
            "Rencana Aksi / Deskripsi": "Penyusunan Standar Operasional Prosedur (SOP) Tata Kelola Kolaboratif Lintas Perangkat Daerah sebagai pedoman kerja bersama.",
            "OPD Lead": "Sekretariat Daerah (Asisten III)",
            "RACI Support": "Bagian Organisasi (R), Bagian Hukum (R), Bapperida (C)",
            "KPI": "Tersedianya dokumen SOP Tata Kelola Kolaboratif.",
            "Timeline": "Agustus 2026 (Mg 4-5)"
        },
        {
            "No": 4,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Sinergi Lintas Sektor",
            "Nama Program / Kegiatan": "Rapat Koordinasi & Sinkronisasi Awal Tim GCG",
            "Rencana Aksi / Deskripsi": "Rapat koordinasi awal dengan Dinas ESDM Provinsi Kalsel & Pengelola Geopark Meratus untuk sinkronisasi kerangka rencana aksi.",
            "OPD Lead": "Bapperida",
            "RACI Support": "Sekretariat Daerah (R), Dinas ESDM Provinsi (C), Geopark Meratus (C)",
            "KPI": "Berita Acara kesepakatan koordinasi & kerangka action plan.",
            "Timeline": "Agustus 2026 (Mg 3-5)"
        },
        {
            "No": 5,
            "Pilar Utama": "Pilar 1: Geodiversity",
            "Aspek Fungsional": "Konservasi",
            "Nama Program / Kegiatan": "Pengenalan awal Geodiversity (Survei Lapangan)",
            "Rencana Aksi / Deskripsi": "Pengenalan lapangan kondisi fisik geosite unggulan di Kotabaru bersama tim agile dan ahli geologi.",
            "OPD Lead": "Bapperida",
            "RACI Support": "Dinas ESDM Prov (R), Bapperida (R)",
            "KPI": "Dokumen hasil survei, koordinat peta geosite, & foto dokumentasi.",
            "Timeline": "September 2026 (Mg 2-3)"
        },
        {
            "No": 6,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Regulasi & Sosialisasi",
            "Nama Program / Kegiatan": "Sosialisasi Potensi Geopark dan Perbup",
            "Rencana Aksi / Deskripsi": "Sosialisasi draf Peraturan Bupati tentang GCG dan potensi Geopark Kotabaru kepada seluruh perangkat daerah dan pemangku kepentingan.",
            "OPD Lead": "Bapperida",
            "RACI Support": "Sekretariat Daerah (R), Disparpora (R), DLH (C)",
            "KPI": "Terselenggaranya pertemuan sosialisasi draf Perbup GCG.",
            "Timeline": "September 2026 (Mg 1-2)"
        },
        {
            "No": 7,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Regulasi dan Kelembagaan",
            "Nama Program / Kegiatan": "Pengesahan & Sosialisasi Perbup GCG",
            "Rencana Aksi / Deskripsi": "Harmonisasi, pengesahan, dan sosialisasi Peraturan Bupati Kotabaru tentang Tata Kelola Kolaboratif Pendukung Geopark.",
            "OPD Lead": "Sekretariat Daerah (Asisten III), Bapperida",
            "RACI Support": "Bagian Hukum (R), Disparpora (R), Bapperida (R), Seluruh OPD (I)",
            "KPI": "Perbup GCG diundangkan dan disosialisasikan ke 100% OPD terkait.",
            "Timeline": "November - Desember 2026"
        },
        {
            "No": 8,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Kelembagaan",
            "Nama Program / Kegiatan": "Pembentukan Tim Koordinasi Geopark Permanen",
            "Rencana Aksi / Deskripsi": "Pembentukan tim koordinasi permanen lintas perangkat daerah melalui Keputusan Bupati untuk mengawal keberlanjutan program.",
            "OPD Lead": "Sekretariat Daerah (Asisten III)",
            "RACI Support": "Bapperida (R), Bagian Organisasi (R), Bagian Hukum (R), BKPSDMD (R)",
            "KPI": "Terbitnya SK Bupati tentang pembentukan Tim Koordinasi Geopark.",
            "Timeline": "November 2026 - Januari 2027"
        },
        {
            "No": 9,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Sistem Informasi",
            "Nama Program / Kegiatan": "Pengembangan Dashboard Monitoring GCG / Forum KOLABORASI",
            "Rencana Aksi / Deskripsi": "Membangun prototype aplikasi Dashboard Monitoring terintegrasi untuk melacak realisasi program kerja fisik & keuangan geopark.",
            "OPD Lead": "Diskominfo",
            "RACI Support": "Bapperida (R), Inspektorat (C), Seluruh OPD (I)",
            "KPI": "Aplikasi dashboard monitoring GCG aktif dan dapat diakses tim.",
            "Timeline": "Januari - Maret 2027"
        },
        {
            "No": 10,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Sistem Informasi",
            "Nama Program / Kegiatan": "Penyebarluasan Informasi",
            "Rencana Aksi / Deskripsi": "Penyebarluasan informasi rutin terkait perkembangan dan pencapaian geopark di Kabupaten Kotabaru kepada masyarakat.",
            "OPD Lead": "Diskominfo",
            "RACI Support": "Bapperida (C), Disparpora (C)",
            "KPI": "Publikasi informasi rutin tentang Geopark di media daerah.",
            "Timeline": "Desember 2026 - Februari 2027"
        },
        {
            "No": 11,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Edukasi",
            "Nama Program / Kegiatan": "FGD I: Integrasi Hasil Kajian Akademis",
            "Rencana Aksi / Deskripsi": "Focus Group Discussion ke-1 untuk mengintegrasikan hasil kajian akademis kebumian Dinas ESDM ke dalam program kegiatan SKPD terkait.",
            "OPD Lead": "Bapperida",
            "RACI Support": "Dinas ESDM Prov (R), Disparpora (R), DLH (C), Tokoh Masyarakat (C)",
            "KPI": "Draf rencana aksi terintegrasi pilar Geodiversity, Biodiversity, & Cultural Diversity.",
            "Timeline": "September 2026 (Mg 4-5)"
        },
        {
            "No": 12,
            "Pilar Utama": "Pilar 1: Geodiversity",
            "Aspek Fungsional": "Edukasi",
            "Nama Program / Kegiatan": "Pembuatan Modul Ajar Muatan Lokal Geopark",
            "Rencana Aksi / Deskripsi": "Menyusun kurikulum muatan lokal sekolah dasar 'Bumi & Alam Saijaan Bersujud' untuk mengedukasi siswa tentang warisan geologi.",
            "OPD Lead": "Disdikbud / BKPSDMD",
            "RACI Support": "Dinas ESDM Prov (C), Pengelola Geopark (C), Sekolah Dasar (I)",
            "KPI": "Draf kurikulum & modul muatan lokal sekolah dasar rampung.",
            "Timeline": "Februari - April 2027"
        },
        {
            "No": 13,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Edukasi & Pelatihan",
            "Nama Program / Kegiatan": "Sosialisasi Publik & Gerakan 'Geopark Goes to School'",
            "Rencana Aksi / Deskripsi": "Kampanye edukasi awal ke sekolah-sekolah di Kotabaru untuk memperkenalkan warisan geologi, konservasi, dan nilai geopark kepada generasi muda.",
            "OPD Lead": "Disdikbud",
            "RACI Support": "Disparpora (R), Dinas ESDM Prov (C), Pengelola Geopark (C), Sekolah (I)",
            "KPI": "Terselenggaranya sosialisasi di minimal 10 sekolah dasar/menengah.",
            "Timeline": "September 2026 (Mg 3-4)"
        },
        {
            "No": 14,
            "Pilar Utama": "Pilar 2: Biodiversity",
            "Aspek Fungsional": "Konservasi",
            "Nama Program / Kegiatan": "Rehabilitasi Hutan Mangrove & Vegetasi Geosite",
            "Rencana Aksi / Deskripsi": "Program penanaman kembali bibit bakau dan tanaman endemik di sekitar pesisir geosite pesisir Kotabaru.",
            "OPD Lead": "DLH",
            "RACI Support": "Dinas Perikanan (R), Komunitas Peduli Lingkungan (R), Swasta/CSR (C)",
            "KPI": "Penanaman minimal 1.000 bibit mangrove dan vegetasi pelindung.",
            "Timeline": "Maret - Juni 2027"
        },
        {
            "No": 15,
            "Pilar Utama": "Pilar 3: Cultural Diversity",
            "Aspek Fungsional": "Kesejahteraan Ekonomi",
            "Nama Program / Kegiatan": "Pembinaan UMKM & Pengembangan Geo-products",
            "Rencana Aksi / Deskripsi": "Pelatihan bagi UMKM lokal untuk meluncurkan produk gastronomi / kuliner khas dan kerajinan tangan bermerek Geopark Kotabaru.",
            "OPD Lead": "Disparpora/Diskoperindag",
            "RACI Support": "Dinas Koperasi & UMKM (R), Komunitas Adat (C), Pelaku Usaha (I)",
            "KPI": "Terbentuknya minimal 5 jenis 'Geo-products' kuliner/gastronomi lokal.",
            "Timeline": "Mei - Juli 2027"
        },
        {
            "No": 16,
            "Pilar Utama": "Pilar 3: Cultural Diversity",
            "Aspek Fungsional": "Kesejahteraan Ekonomi",
            "Nama Program / Kegiatan": "Penyelenggaraan Festival Budaya & Ekowisata Geopark Kotabaru",
            "Rencana Aksi / Deskripsi": "Mengintegrasikan promosi warisan geologi dalam pameran khusus geosite di sela-sela event budaya tahunan (seperti Festival Budaya Saijaan) guna menarik wisatawan.",
            "OPD Lead": "Disparpora",
            "RACI Support": "Disdikbud (R), DLH (C), Komunitas Seni/Budaya (R), Pelaku Usaha (I)",
            "KPI": "Terlaksananya stand khusus pameran warisan geologi dan pertunjukan budaya bertema bumi.",
            "Timeline": "Maret - April 2027"
        },
        {
            "No": 17,
            "Pilar Utama": "Pilar 1: Geodiversity",
            "Aspek Fungsional": "Edukasi & Pelatihan",
            "Nama Program / Kegiatan": "Pelatihan & Sertifikasi Pemandu Wisata Geologi (Geo-Guides)",
            "Rencana Aksi / Deskripsi": "Pelatihan teknis bagi pemuda lokal and Pokdarwis mengenai sejarah geologi Kotabaru, teknik pemanduan, dan konservasi alam.",
            "OPD Lead": "Disparpora",
            "RACI Support": "Dinas ESDM Prov (R), DLH (C), HPI Kotabaru (R), BKPSDMD (C), Komunitas Lokal (I)",
            "KPI": "Lahirnya minimal 15 pemandu wisata lokal bersertifikat khusus Geopark.",
            "Timeline": "Mei - Juni 2027"
        },
        {
            "No": 18,
            "Pilar Utama": "Pilar 3: Cultural Diversity",
            "Aspek Fungsional": "Kesejahteraan Ekonomi",
            "Nama Program / Kegiatan": "Seminar Klaster Saijaan-Bersujud",
            "Rencana Aksi / Deskripsi": "Menyelenggarakan seminar ilmiah nasional untuk mengenalkan potensi klaster Saijaan-Bersujud kepada publik dan menarik minat akademisi.",
            "OPD Lead": "Bapperida",
            "RACI Support": "Dinas ESDM Prov (C), DPMD (R), Pelaku Usaha (I)",
            "KPI": "Terselenggaranya seminar klaster Saijaan-Bersujud.",
            "Timeline": "Maret 2027"
        },
        {
            "No": 19,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Sinergi Lintas Sektor",
            "Nama Program / Kegiatan": "Sinkronisasi Tata Ruang",
            "Rencana Aksi / Deskripsi": "Sinkronisasi dan harmonisasi pemanfaatan ruang di Kabupaten Kotabaru agar sesuai dengan zonasi konservasi Geopark.",
            "OPD Lead": "PUPR",
            "RACI Support": "Bapperida (R), DLH (C)",
            "KPI": "Terlaksananya sinkronisasi tata ruang wilayah Geopark.",
            "Timeline": "Juni - Juli 2027"
        },
        {
            "No": 20,
            "Pilar Utama": "Pilar 1: Geodiversity",
            "Aspek Fungsional": "Kesejahteraan Ekonomi",
            "Nama Program / Kegiatan": "Penyusunan Peta Wisata Geologi & Media Informasi Geosite",
            "Rencana Aksi / Deskripsi": "Inisiasi pembuatan brosur digital, peta rute geotourism, dan sistem QR code informasi sejarah geologi pada geosite unggulan Kotabaru.",
            "OPD Lead": "Disparpora",
            "RACI Support": "Dinas ESDM Prov (R), Diskominfo (R)",
            "KPI": "Rilis peta wisata digital geopark Kotabaru dan pemasangan QR code info di 3 geosite utama.",
            "Timeline": "September 2026 (Mg 4) - Oktober 2026 (Mg 1)"
        },
        {
            "No": 21,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Kelembagaan",
            "Nama Program / Kegiatan": "Pelembagaan GCG ke Dokumen Perencanaan Makro",
            "Rencana Aksi / Deskripsi": "Mengintegrasikan indikator kinerja geopark & rencana aksi kolaboratif ke dalam dokumen RPJMD, Renstra OPD, dan RKPD.",
            "OPD Lead": "Bapperida",
            "RACI Support": "Sekretariat Daerah (Asisten III) (R), Seluruh OPD (R)",
            "KPI": "Indikator GCG termuat dalam RKPD & APBD Kabupaten Kotabaru.",
            "Timeline": "Juli - Agustus 2027"
        },
        {
            "No": 22,
            "Pilar Utama": "Enabler / Tata Kelola",
            "Aspek Fungsional": "Sistem Informasi",
            "Nama Program / Kegiatan": "Implementasi Integrasi One Data Geopark",
            "Rencana Aksi / Deskripsi": "Menyatukan data spasial, data keanekaragaman hayati, data cagar budaya, dan statistik kunjungan ke platform One Data Kotabaru.",
            "OPD Lead": "Diskominfo",
            "RACI Support": "Bapperida (R), Disparpora (R), Disdikbud (R), DLH (R)",
            "KPI": "Portal data interaktif 'One Data Geopark' rilis publik.",
            "Timeline": "September - Oktober 2027"
        },
        {
            "No": 23,
            "Pilar Utama": "Pilar 3: Cultural Diversity",
            "Aspek Fungsional": "Kesejahteraan Ekonomi",
            "Nama Program / Kegiatan": "Promosi & Sertifikasi Ekowisata Internasional",
            "Rencana Aksi / Deskripsi": "Mendaftarkan ekowisata Kotabaru untuk meraih sertifikasi pariwisata berkelanjutan dan mempromosikannya ke jejaring global.",
            "OPD Lead": "Disparpora",
            "RACI Support": "ASCG / GGN Network (C), Agen Travel (I), Media Massa (I)",
            "KPI": "Geopark Kotabaru masuk ke dalam kalender wisata nasional/global.",
            "Timeline": "November 2027 - Seterusnya"
        }
    ]

# Sidebar Navigation
st.sidebar.image("https://img.icons8.com/color/96/000000/earth-exploration.png", width=80)
st.sidebar.title("FGD I Geopark Kotabaru (Road Map V3)")
st.sidebar.markdown("**Integrasi Peta Jalan (Road Map) V3**\n*Kabupaten Kotabaru - 23 September 2026*")
st.sidebar.markdown("---")

menu = st.sidebar.radio(
    "Navigasi Instrumen:",
    ["🏠 Beranda & Panduan", "⏱️ Timer & Rundown", "👥 Desk Diskusi Interaktif", "📊 Hasil Sinkronisasi Matriks", "✍️ Draf Berita Acara"]
)

st.sidebar.markdown("---")
st.sidebar.info(
    "💡 **Petunjuk Deploy:**\n\n"
    "Aplikasi ini berbasis **Streamlit (Python)**. Anda dapat langsung menjalankan secara lokal dengan perintah:\n"
    "`streamlit run fgd_app.py`\n\n"
    "Atau mendeploy secara gratis ke **Streamlit Community Cloud** atau **Hugging Face Spaces**."
)

# 1. HOME SHEET
if menu == "🏠 Beranda & Panduan":
    st.title("🌍 Beranda & Panduan FGD I Geopark Kotabaru (Road Map V3)")
    st.subheader("Integrasi Peta Jalan (Road Map) v3 Baru")
    
    st.markdown(
        """
        ### 📌 Latar Belakang & Tujuan FGD I
        Fokus utama forum ini adalah menyelaraskan langkah strategis berdasarkan **Road Map v3** ke dalam pilar taktis pengembangan fisik, konservasi, edukasi, dan ekonomi lokal di Kabupaten Kotabaru.
        Sesi pleno diawali dengan paparan dari **Dinas ESDM Provinsi Kalimantan Selatan** selaku narasumber utama untuk memberikan stimulasi kebijakan geopark di Kalsel, dilanjutkan ekspose Road Map v3, diskusi kelompok (FGD) dengan instrumen, dan penandatanganan Berita Acara.
        
        ### 🛠️ Panduan Penggunaan Aplikasi Interaktif
        Aplikasi ini berfungsi sebagai **instrumen digital pendamping** selama diskusi meja kelompok terfokus (*breakout desk*). 
        1. Gunakan tab **⏱️ Timer & Rundown** untuk memandu jalannya sesi agar tepat waktu.
        2. Masuk ke tab **👥 Desk Diskusi Interaktif** saat sesi kelompok dimulai. Pilih pilar desk Anda, gunakan pertanyaan pemantik untuk memandu debat, dan catat usulan program baru langsung ke dalam form.
        3. Lihat, edit, dan unduh kompilasi rencana aksi secara langsung pada tab **📊 Hasil Sinkronisasi Matriks** sebagai format Excel/CSV yang siap pakai.
        4. Generate dokumen naskah penandatanganan kesepakatan secara instan pada tab **✍️ Draf Berita Acara**.
        """
    )
    
    # Quick KPI cards
    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Tanggal FGD", "23 September 2026")
    with col2:
        st.metric("Total Program Road Map V3", f"{len(st.session_state.matrix_data)} Kegiatan")
    with col3:
        st.metric("Mitra Narasumber", "Dinas ESDM Prov Kalsel")

# 2. TIMER & RUNDOWN SHEET
elif menu == "⏱️ Timer & Rundown":
    st.title("⏱️ Manajemen Waktu & Rundown FGD I")
    
    # Countdown timer simulation
    st.subheader("⏱️ Pengendali Waktu Sesi Rapat")
    st.info(f"Target Penyelesaian Forum: **Rabu, 23 September 2026 - 16:00 WITA**")
    
    st.markdown("### 📋 Susunan Acara (Rundown)")
    rundown_data = [
        {"Waktu (WITA)": "08.00 - 08.30", "Agenda": "Registrasi Peserta & Pembagian Berkas (Instrumen FGD)", "PJ": "Panitia Sekretariat"},
        {"Waktu (WITA)": "08.30 - 09.00", "Agenda": "Pembukaan, Menyanyikan Lagu Kebangsaan, & Sambutan Bupati/Sekda", "PJ": "Bupati Kotabaru / Sekda"},
        {"Waktu (WITA)": "09.00 - 10.30", "Agenda": "Sesi I: Stimulasi Geopark & Kebijakan Pengembangan Geopark di Kalsel", "PJ": "Dinas ESDM Provinsi Kalsel (Narasumber)"},
        {"Waktu (WITA)": "10.30 - 11.30", "Agenda": "Sesi II: Ekspose Road Map Pengembangan Geopark Kotabaru V3", "PJ": "Bapperida & Disparpora"},
        {"Waktu (WITA)": "11.30 - 12.00", "Agenda": "Penjelasan Teknis & Pembagian Desk Diskusi Kelompok", "PJ": "Sekretariat Geopark"},
        {"Waktu (WITA)": "12.00 - 13.00", "Agenda": "Istirahat, Sholat, Makan Siang (ISOMA)", "PJ": "Seluruh Peserta"},
        {"Waktu (WITA)": "13.00 - 15.00", "Agenda": "Sesi Breakout Diskusi Kelompok Terfokus (3 Desk Pilar) dengan Instrumen", "PJ": "Moderator Desk & OPD Teknis"},
        {"Waktu (WITA)": "15.00 - 15.30", "Agenda": "Presentasi Hasil Rumusan Setiap Desk", "PJ": "Perwakilan Kelompok"},
        {"Waktu (WITA)": "15.30 - 16.00", "Agenda": "Penandatanganan Berita Acara Komitmen Bersama & Penutupan", "PJ": "Pimpinan OPD, Dinas ESDM, & BKPSDMD"}
    ]
    st.table(pd.DataFrame(rundown_data))

# 3. INTERACTIVE DESK DISCUSSION
elif menu == "👥 Desk Diskusi Interaktif":
    st.title("👥 Sesi Breakout Desk Diskusi Kelompok")
    st.markdown("Gunakan instrumen interaktif ini pada masing-masing meja kelompok pilar untuk mengarahkan diskusi dan mencatat hasil kesepakatan secara instan.")

    desk_choice = st.selectbox(
        "Pilih Meja Diskusi (Desk Pilar) Anda:",
        [
            "Desk 1: Pilar Keragaman Geologi (Geodiversity)", 
            "Desk 2: Pilar Keanekaragaman Hayati (Biodiversity)", 
            "Desk 3: Pilar Keragaman Budaya (Cultural Diversity)"
        ]
    )

    if desk_choice == "Desk 1: Pilar Keragaman Geologi (Geodiversity)":
        st.subheader("📐 Desk 1: Pilar Keragaman Geologi (Geodiversity)")
        st.markdown("**Fokus Pokok:** Konservasi warisan geologi, pengembangan pariwisata minat khusus (*geotourism*), dan edukasi kebumian.")
        
        # Trigger Questions Expander
        with st.expander("❓ Pertanyaan Pemantik Diskusi (Trigger Questions) - Klik untuk membuka"):
            st.markdown(
                """
                1. **Prioritas Geosite & KCAG:** Mana saja dari geosite hasil kajian Dinas ESDM yang mendesak untuk segera dipasangi papan batas konservasi fisik dalam jangka pendek?
                2. **Pemandu Lokal & BKPSDMD:** Bagaimana sinergi Disparpora dan BKPSDMD untuk sertifikasi kompetensi pemandu geowisata (*Geo-guides*) bagi pemuda lokal?
                3. **Muatan Lokal:** Bagaimana integrasi kurikulum muatan lokal sekolah dasar "Bumi & Alam Saijaan Bersujud" dengan dukungan penataran guru oleh Disdikbud dan BKPSDMD?
                """
            )
            
    elif desk_choice == "Desk 2: Pilar Keanekaragaman Hayati (Biodiversity)":
        st.subheader("🌱 Desk 2: Pilar Keanekaragaman Hayati (Biodiversity)")
        st.markdown("**Fokus Pokok:** Perlindungan flora dan fauna khas daerah, konservasi ekosistem pesisir, serta pelestarian hutan bakau (mangrove) Kotabaru.")
        
        with st.expander("❓ Pertanyaan Pemantik Diskusi (Trigger Questions) - Klik untuk membuka"):
            st.markdown(
                """
                1. **Rehabilitasi Ekosistem:** OPD mana saja yang bersinergi langsung (misalnya DLH dan Dinas Perikanan) untuk melakukan penanaman vegetasi mangrove pelindung di sekitar geosite pesisir?
                2. **Pelibatan Komunitas:** Bagaimana cara efektif menggerakkan Kelompok Masyarakat Pengawas (Pokmaswas) dan komunitas peduli lingkungan lokal agar ikut serta menjaga ekosistem penunjang geopark?
                """
            )

    else:
        st.subheader("🎭 Desk 3: Pilar Keragaman Budaya (Cultural Diversity)")
        st.markdown("**Fokus Pokok:** Perlindungan cagar budaya, dokumentasi adat istiadat pesisir, serta promosi produk kreatif berbasis gastronomi bumi (*Geo-products*).")
        
        with st.expander("❓ Pertanyaan Pemantik Diskusi (Trigger Questions) - Klik untuk membuka"):
            st.markdown(
                """
                1. **Geo-products & Kuliner (Jangka Panjang):** Kuliner atau kerajinan tangan lokal khas Kotabaru apa yang paling potensial diberikan branding visual dan kurasi sebagai *Geo-products* resmi dalam fase jangka panjang?
                2. **Sinergi Event Budaya:** Bagaimana detail integrasi pameran edukasi bumi dan sains geologi ke dalam agenda rutin Festival Budaya Saijaan bersama Disparpora dan Disdikbud?
                """
            )

    # Interactive Form to Add Program to Matrix
    st.markdown("---")
    st.subheader("➕ Form Input Usulan Program Rencana Aksi Baru / Hasil Penyesuaian")
    
    with st.form("add_program_form"):
        col1, col2 = st.columns(2)
        with col1:
            pilar_opt = st.selectbox("Pilar Utama:", ["Pilar 1: Geodiversity", "Pilar 2: Biodiversity", "Pilar 3: Cultural Diversity", "Enabler / Tata Kelola"])
            aspek_opt = st.selectbox("Aspek Fungsional:", ["Konservasi", "Edukasi & Pelatihan", "Kesejahteraan Ekonomi", "Sistem Informasi", "Kelembagaan & Regulasi"])
            prog_name = st.text_input("Nama Program / Kegiatan Kerja:", placeholder="Contoh: Pembuatan Peta Jalur Geowisata...")
            timeline_opt = st.text_input("Estimasi Jadwal (Timeline Detail):", placeholder="Contoh: September - Oktober 2026")
        
        with col2:
            opd_lead = st.selectbox("OPD Penanggung Jawab (Lead):", ["Bapperida", "Disparpora", "Disdikbud", "DLH", "Diskominfo", "Dinas PUPR", "Sekretariat Daerah (Asisten III)", "BKPSDMD"])
            raci_supp = st.text_input("OPD Pendukung & Pembagian Peran RACI:", placeholder="Contoh: BKPSDMD (R), Dinas ESDM Prov (R), Bagian Hukum (R)")
            kpi_desc = st.text_area("Indikator Keberhasilan (KPI):", placeholder="Contoh: Tersedianya modul kurikulum muatan lokal di tingkat SD...")
            prog_desc = st.text_area("Deskripsi Ringkas Kegiatan:", placeholder="Jelaskan mekanisme pelaksanaan kegiatan secara kolaboratif...")
            
        submitted = st.form_submit_button("💾 Masukkan ke dalam Draft Matriks Sementara")
        if submitted:
            if prog_name and prog_desc:
                new_item = {
                    "No": len(st.session_state.matrix_data) + 1,
                    "Pilar Utama": pilar_opt,
                    "Aspek Fungsional": aspek_opt,
                    "Nama Program / Kegiatan": prog_name,
                    "Rencana Aksi / Deskripsi": prog_desc,
                    "OPD Lead": opd_lead,
                    "RACI Support": raci_supp,
                    "KPI": kpi_desc,
                    "Timeline": timeline_opt
                }
                st.session_state.matrix_data.append(new_item)
                st.success(f"✔️ Program '{prog_name}' berhasil ditambahkan ke draf matriks sementara!")
            else:
                st.error("⚠️ Nama program dan deskripsi ringkas wajib diisi sebelum menyimpan.")

# 4. VIEW & EXPORT MATRIX
elif menu == "📊 Hasil Sinkronisasi Matriks":
    st.title("📊 Hasil Sinkronisasi Matriks Road Map v3")
    st.markdown("Berikut adalah tabel dinamis draf Road Map Geopark Kotabaru yang berhasil dikompilasi lintas meja desk.")
    
    # Display table in DataFrame format
    df = pd.DataFrame(st.session_state.matrix_data)
    st.dataframe(df, use_container_width=True)
    
    # Download as CSV Option
    csv_data = df.to_csv(index=False).encode('utf-8')
    st.download_button(
        label="📥 Download Matriks Road Map Hasil FGD (Format CSV/Excel Compatible)",
        data=csv_data,
        file_name="matriks_roadmap_geopark_fgd_v3.csv",
        mime="text/csv"
    )
    
    # Chart or summary representation
    st.subheader("📈 Ringkasan Kontribusi OPD (Lead)")
    opd_counts = df["OPD Lead"].value_counts()
    st.bar_chart(opd_counts)

# 5. DRAFT BERITA ACARA
else:
    st.title("✍️ Penandatanganan Berita Acara Kesepakatan FGD I")
    st.markdown("Di akhir pelaksanaan FGD, draf naskah kesepakatan komitmen bersama ini dicetak dan ditandatangani oleh perwakilan pimpinan OPD kunci sebagai bukti kesepakatan program.")
    
    # Berita Acara content
    berita_acara_title = "BERITA ACARA KESEPAKATAN HASIL FGD I"
    berita_acara_body = (
        "Pada hari ini, Rabu tanggal Dua Puluh Tiga bulan September tahun Dua Ribu Dua Puluh Enam (23-09-2026), "
        "bertempat di Kabupaten Kotabaru, seluruh unsur pemangku kepentingan (Pentahelix) pengembangan Geopark Kotabaru "
        "(Klaster Saijaan Bersujud) yang bertandatangan di bawah ini telah menyepakati draf Road Map Pengembangan "
        "Geopark Kotabaru V3 Jangka Pendek, Menengah, dan Panjang berbasis Tiga Pilar Geopark "
        "(Geodiversity, Biodiversity, Cultural Diversity)."
    )
    kesepakatan_points = [
        "Menyetujui hasil delineasi awal inventarisasi geosite ilmiah hasil kerja sama dengan Dinas ESDM Provinsi Kalsel guna mengusulkan Kawasan Cagar Alam Geologi (KCAG) ke Kementerian ESDM.",
        "Berkomitmen menyinkronkan dan mengintegrasikan program kerja sektoral masing-masing OPD/SKPD, termasuk program pembinaan aparatur daerah oleh BKPSDMD ke dalam Road Map terpadu.",
        "Mendukung penuh penyusunan kurikulum muatan lokal sekolah dasar 'Bumi & Alam Saijaan Bersujud' serta menyepakati pergeseran program pembinaan UMKM & Geo-products ke jangka panjang demi kesiapan infrastruktur daerah."
    ]
    penutup = "Demikian Berita Acara ini dibuat dengan kesadaran penuh demi kelestarian bumi dan kesejahteraan masyarakat Kotabaru."
    
    pihak_pemkab = [
        "Asisten III Administrasi Umum Setda",
        "Kepala Bapperida Kotabaru",
        "Kepala Disparpora Kotabaru",
        "Kepala BKPSDMD Kotabaru"
    ]
    pihak_mitra = [
        "Dinas ESDM Provinsi Kalimantan Selatan",
        "Badan Pengelola Geopark Meratus"
    ]
    
    with st.container():
        st.markdown(f"### **{berita_acara_title}**")
        st.markdown(
            f"""
            {berita_acara_body}
            
            **Poin-Poin Kesepakatan Utama:**
            1. {kesepakatan_points[0]}
            2. {kesepakatan_points[1]}
            3. {kesepakatan_points[2]}
            
            {penutup}
            """
        )
        
        st.markdown("---")
        st.markdown("#### **Daftar Pihak yang Menyetujui:**")
        col1, col2 = st.columns(2)
        with col1:
            st.write("✒️ **Pemerintah Kabupaten Kotabaru**")
            for p in pihak_pemkab:
                st.write(f"- {p}")
        with col2:
            st.write("✒️ **Mitra Jaringan & Provinsi**")
            for m in pihak_mitra:
                st.write(f"- {m}")
    
    # --- Export Options ---
    st.markdown("---")
    st.subheader("📄 Ekspor Dokumen Berita Acara")
    
    include_matrix = st.checkbox("Sertakan tabel Matriks Road Map dalam dokumen ekspor", value=True)
    
    col_pdf, col_docx, col_print = st.columns(3)
    
    # === PDF Export ===
    with col_pdf:
        def generate_pdf():
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=25)
            
            # --- Page 1: Berita Acara (Portrait, legal margins) ---
            pdf.add_page()
            pdf.set_margins(left=30, top=25, right=25)  # Standard legal: 3cm left, 2.5cm top/right
            pdf.set_y(25)
            
            # Title
            pdf.set_font("Helvetica", "B", 14)
            pdf.multi_cell(0, 8, berita_acara_title, align="C")
            pdf.ln(3)
            
            # Underline
            pdf.set_draw_color(0, 0, 0)
            pdf.set_line_width(0.5)
            x_start = pdf.get_x()
            y_line = pdf.get_y()
            pdf.line(x_start + 20, y_line, x_start + 135, y_line)
            pdf.ln(8)
            
            # Body
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, berita_acara_body, align="J")
            pdf.ln(6)
            
            # Kesepakatan points
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 8, "Poin-Poin Kesepakatan Utama:", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            pdf.set_font("Helvetica", "", 11)
            for i, point in enumerate(kesepakatan_points, 1):
                pdf.multi_cell(0, 6, f"{i}. {point}", align="J")
                pdf.ln(3)
            
            pdf.ln(4)
            pdf.multi_cell(0, 6, penutup, align="J")
            pdf.ln(8)
            
            # Signatories section - keep together on same page
            # Estimate space needed: ~100mm for full sign block
            sign_block_height = 100
            if pdf.get_y() + sign_block_height > 272:  # 297 - 25 bottom margin
                pdf.add_page()
                pdf.set_margins(left=30, top=25, right=25)
                pdf.set_y(25)
            
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 8, "Daftar Pihak yang Menyetujui:", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(3)
            
            usable_width = 210 - 30 - 25  # page width - left margin - right margin = 155mm
            half_w = usable_width / 2
            
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(half_w, 7, "Pemerintah Kabupaten Kotabaru:", new_x="END", new_y="TOP")
            pdf.cell(half_w, 7, "Mitra Jaringan & Provinsi:", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
            pdf.set_font("Helvetica", "", 10)
            
            max_rows = max(len(pihak_pemkab), len(pihak_mitra))
            for i in range(max_rows):
                left = f"- {pihak_pemkab[i]}" if i < len(pihak_pemkab) else ""
                right = f"- {pihak_mitra[i]}" if i < len(pihak_mitra) else ""
                pdf.cell(half_w, 6, left, new_x="END", new_y="TOP")
                pdf.cell(half_w, 6, right, new_x="LMARGIN", new_y="NEXT")
            
            # Signature block
            pdf.ln(12)
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(half_w, 6, "", new_x="END", new_y="TOP")
            pdf.cell(half_w, 6, "Kotabaru, 23 September 2026", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(18)
            pdf.cell(half_w, 6, "", new_x="END", new_y="TOP")
            pdf.cell(half_w, 6, "(___________________________)", new_x="LMARGIN", new_y="NEXT")
            
            # --- Page 2+: Matrix Table (Landscape for more space) ---
            if include_matrix:
                pdf.add_page(orientation="L")
                pdf.set_margins(left=15, top=15, right=15)
                pdf.set_auto_page_break(auto=True, margin=15)
                pdf.set_y(15)
                
                pdf.set_font("Helvetica", "B", 12)
                pdf.cell(0, 10, "LAMPIRAN: Matriks Road Map Geopark Kotabaru V3", new_x="LMARGIN", new_y="NEXT", align="C")
                pdf.ln(5)
                
                df = pd.DataFrame(st.session_state.matrix_data)
                pdf_cols = ["No", "Pilar Utama", "Aspek Fungsional", "Nama Program / Kegiatan", "OPD Lead", "KPI", "Timeline"]
                # Landscape A4 usable: 297 - 15 - 15 = 267mm
                col_widths = [8, 35, 30, 60, 35, 55, 44]  # total = 267mm
                
                line_h = 5
                font_size = 7
                
                # Table header
                pdf.set_font("Helvetica", "B", font_size)
                pdf.set_fill_color(220, 220, 220)
                for i, col_name in enumerate(pdf_cols):
                    pdf.cell(col_widths[i], line_h + 2, col_name, border=1, align="C", fill=True)
                pdf.ln()
                
                # Table body with text wrapping
                pdf.set_font("Helvetica", "", font_size)
                for _, row in df.iterrows():
                    # Calculate max lines needed for this row
                    max_lines = 1
                    cell_texts = []
                    for i, col_name in enumerate(pdf_cols):
                        text = str(row[col_name])
                        cell_texts.append(text)
                        # Estimate lines needed
                        char_per_line = int(col_widths[i] / (font_size * 0.5))
                        lines_needed = max(1, -(-len(text) // char_per_line))  # ceiling division
                        max_lines = max(max_lines, lines_needed)
                    
                    row_height = line_h * max_lines
                    
                    # Check page break
                    if pdf.get_y() + row_height > 195:  # landscape height - bottom margin
                        pdf.add_page(orientation="L")
                        pdf.set_margins(left=15, top=15, right=15)
                        pdf.set_y(15)
                        # Reprint header
                        pdf.set_font("Helvetica", "B", font_size)
                        pdf.set_fill_color(220, 220, 220)
                        for i, col_name in enumerate(pdf_cols):
                            pdf.cell(col_widths[i], line_h + 2, col_name, border=1, align="C", fill=True)
                        pdf.ln()
                        pdf.set_font("Helvetica", "", font_size)
                    
                    # Draw cells with multi_cell for wrapping
                    x_start = pdf.get_x()
                    y_start = pdf.get_y()
                    
                    for i, text in enumerate(cell_texts):
                        pdf.set_xy(x_start + sum(col_widths[:i]), y_start)
                        # Draw border rectangle
                        pdf.rect(x_start + sum(col_widths[:i]), y_start, col_widths[i], row_height)
                        # Draw text inside with padding
                        pdf.set_xy(x_start + sum(col_widths[:i]) + 1, y_start + 1)
                        pdf.multi_cell(col_widths[i] - 2, line_h, text, align="L")
                    
                    pdf.set_y(y_start + row_height)
            
            return bytes(pdf.output())
        
        pdf_bytes = generate_pdf()
        st.download_button(
            label="📥 Download PDF",
            data=pdf_bytes,
            file_name="Berita_Acara_FGD_I_Geopark_Kotabaru.pdf",
            mime="application/pdf",
            use_container_width=True
        )
    
    # === DOCX Export ===
    with col_docx:
        def generate_docx():
            doc = Document()
            
            # Set margins
            for section in doc.sections:
                section.top_margin = Cm(2.5)
                section.bottom_margin = Cm(2.5)
                section.left_margin = Cm(3)
                section.right_margin = Cm(2.5)
            
            # Title
            title = doc.add_heading(berita_acara_title, level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Body
            p = doc.add_paragraph(berita_acara_body)
            p.paragraph_format.space_after = Pt(12)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            # Kesepakatan points
            doc.add_paragraph("Poin-Poin Kesepakatan Utama:", style="Heading 3")
            for point in kesepakatan_points:
                p = doc.add_paragraph(point, style="List Number")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            doc.add_paragraph()
            p = doc.add_paragraph(penutup)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            # Signatories
            doc.add_paragraph()
            doc.add_heading("Daftar Pihak yang Menyetujui:", level=3)
            
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            
            # Header row
            hdr = table.rows[0].cells
            hdr[0].text = "Pemerintah Kabupaten Kotabaru"
            hdr[1].text = "Mitra Jaringan & Provinsi"
            for cell in hdr:
                for paragraph in cell.paragraphs:
                    paragraph.runs[0].bold = True
            
            # Content rows
            max_rows = max(len(pihak_pemkab), len(pihak_mitra))
            for i in range(max_rows):
                row = table.add_row().cells
                row[0].text = f"- {pihak_pemkab[i]}" if i < len(pihak_pemkab) else ""
                row[1].text = f"- {pihak_mitra[i]}" if i < len(pihak_mitra) else ""
            
            # Signature area
            doc.add_paragraph()
            doc.add_paragraph()
            sig_table = doc.add_table(rows=3, cols=2)
            sig_table.rows[0].cells[0].text = "Kotabaru, 23 September 2026"
            sig_table.rows[2].cells[0].text = "(___________________________)"
            sig_table.rows[2].cells[1].text = "(___________________________)"
            
            # Include matrix if requested
            if include_matrix:
                doc.add_page_break()
                doc.add_heading("LAMPIRAN: Matriks Road Map Geopark Kotabaru V3", level=2)
                
                df = pd.DataFrame(st.session_state.matrix_data)
                pdf_cols = ["No", "Pilar Utama", "Nama Program / Kegiatan", "OPD Lead", "KPI", "Timeline"]
                
                matrix_table = doc.add_table(rows=1, cols=len(pdf_cols))
                matrix_table.style = "Table Grid"
                
                # Header
                for i, col_name in enumerate(pdf_cols):
                    matrix_table.rows[0].cells[i].text = col_name
                    for paragraph in matrix_table.rows[0].cells[i].paragraphs:
                        paragraph.runs[0].bold = True
                
                # Data rows
                for _, row_data in df.iterrows():
                    row = matrix_table.add_row().cells
                    for i, col_name in enumerate(pdf_cols):
                        row[i].text = str(row_data[col_name])
                
                # Set font size for table
                for row in matrix_table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.style.font.size = Pt(8)
            
            # Save to bytes
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()
        
        docx_bytes = generate_docx()
        st.download_button(
            label="📥 Download Word (.docx)",
            data=docx_bytes,
            file_name="Berita_Acara_FGD_I_Geopark_Kotabaru.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    
    # === Print (Browser) ===
    with col_print:
        st.markdown(
            """
            <style>
            @media print {
                .stApp { background: white; }
                .stSidebar, .stButton, .stDownloadButton, header, footer { display: none !important; }
            }
            .print-btn-custom {
                width: 100%;
                padding: 0.5rem 1rem;
                background-color: #FF4B4B;
                color: white;
                border: none;
                border-radius: 0.5rem;
                cursor: pointer;
                font-size: 0.875rem;
                font-weight: 600;
                font-family: "Source Sans Pro", sans-serif;
                height: 2.5rem;
                display: flex;
                align-items: center;
                justify-content: center;
                gap: 0.25rem;
                transition: background-color 0.2s;
            }
            .print-btn-custom:hover {
                background-color: #FF2B2B;
            }
            </style>
            <button class="print-btn-custom" onclick="window.print()">🖨️ Cetak / Print (Browser)</button>
            """,
            unsafe_allow_html=True
        )
